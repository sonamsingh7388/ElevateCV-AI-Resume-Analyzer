"""
ATS Resume Generator & DOCX Export Engine.
Generates optimized, ATS-compliant resumes with upgraded action verbs,
injected target keywords, and exports them to clean Microsoft Word (.docx) documents.
"""

import io
import re
from typing import Dict, List, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from modules.suggestions import ACTION_VERB_REPLACEMENTS


def optimize_sections(
    sections: Dict[str, str],
    missing_skills: Optional[List[str]] = None,
    weak_verbs: Optional[List[str]] = None
) -> Dict[str, str]:
    """
    Auto-optimizes resume sections:
    - Replaces passive phrases with high-impact action verbs.
    - Integrates missing target skills into the Skills section.
    - Standardizes bullet formatting.
    """
    optimized: Dict[str, str] = {}
    missing_skills = missing_skills or []

    # Map for easy replacement
    replacements = {
        r"\bworked on\b": "Engineered",
        r"\bresponsible for\b": "Orchestrated",
        r"\bhelped with\b": "Collaborated to deliver",
        r"\bassisted in\b": "Partnered to implement",
        r"\bhandled\b": "Directed",
        r"\btasked with\b": "Spearheaded",
        r"\bduties included\b": "Executed"
    }

    for sec_name, content in sections.items():
        if not content:
            continue
        upgraded_content = content

        # Replace weak verbs
        for pattern, replacement in replacements.items():
            upgraded_content = re.sub(pattern, replacement, upgraded_content, flags=re.IGNORECASE)

        # Standardize bullets: replace weird unicode or dashes with bullet character
        lines = upgraded_content.split("\n")
        cleaned_lines = []
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            # If line starts with -, *, or strange bullet, clean it
            if stripped.startswith(("-", "*", "•", "–", "—", ">")):
                stripped_bullet = re.sub(r"^[-*•–—>]\s*", "", stripped)
                cleaned_lines.append(f"• {stripped_bullet}")
            else:
                cleaned_lines.append(stripped)

        # If this is the skills section, append missing target skills if any
        if sec_name.lower() == "skills" and missing_skills:
            present_text = "\n".join(cleaned_lines).lower()
            unadded = [s for s in missing_skills if s.lower() not in present_text]
            if unadded:
                top_to_add = unadded[:6]
                cleaned_lines.append(f"• Targeted Competencies: {', '.join(top_to_add)}")

        optimized[sec_name] = "\n".join(cleaned_lines)

    return optimized


def generate_full_resume_markdown(
    candidate_name: str,
    contact_info: Dict[str, Any],
    sections_dict: Dict[str, str]
) -> str:
    """Creates a markdown representation of the optimized resume."""
    md = []
    md.append(f"# {candidate_name.upper()}\n")

    # Contact line
    contacts = []
    if contact_info.get("email"):
        contacts.append(contact_info["email"])
    if contact_info.get("phone"):
        contacts.append(contact_info["phone"])
    if contact_info.get("linkedin"):
        contacts.append(contact_info["linkedin"])
    if contact_info.get("github"):
        contacts.append(contact_info["github"])

    if contacts:
        md.append(" | ".join(contacts))
        md.append("\n---\n")

    # Order of standard sections
    section_order = ["summary", "skills", "experience", "projects", "education", "certifications", "awards"]
    seen_sections = set()

    for sec in section_order:
        if sec in sections_dict and sections_dict[sec].strip():
            seen_sections.add(sec)
            md.append(f"## {sec.upper()}\n")
            md.append(sections_dict[sec].strip())
            md.append("\n")

    # Any remaining custom sections
    for sec, content in sections_dict.items():
        if sec not in seen_sections and content.strip() and sec != "header":
            md.append(f"## {sec.upper()}\n")
            md.append(content.strip())
            md.append("\n")

    return "\n".join(md)


def _set_cell_border(cell, **kwargs):
    """Utility to set cell borders in python-docx."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            for key in ('val', 'color', 'sz', 'space'):
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def export_resume_to_docx(
    candidate_name: str,
    contact_info: Dict[str, Any],
    sections_dict: Dict[str, str]
) -> bytes:
    """
    Builds an executive, ATS-compliant Microsoft Word document (.docx).
    Uses clean typography (Calibri), 0.75-inch margins, clear section headings,
    and bullet indentations.
    """
    doc = docx.Document()

    # Configure 0.75-inch page margins (ATS Gold Standard)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Style: Normal
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(10)
    font.color.rgb = RGBColor(30, 41, 59)

    # 1. Candidate Name (Header)
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run(candidate_name.upper())
    run_name.font.name = 'Calibri'
    run_name.font.size = Pt(18)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(15, 23, 42)

    # 2. Contact Information
    contact_parts = []
    if contact_info.get("email"):
        contact_parts.append(contact_info["email"])
    if contact_info.get("phone"):
        contact_parts.append(contact_info["phone"])
    if contact_info.get("linkedin"):
        contact_parts.append(contact_info["linkedin"].replace("https://", "").replace("http://", ""))
    if contact_info.get("github"):
        contact_parts.append(contact_info["github"].replace("https://", "").replace("http://", ""))

    if contact_parts:
        p_contact = doc.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_contact.paragraph_format.space_before = Pt(0)
        p_contact.paragraph_format.space_after = Pt(10)
        run_contact = p_contact.add_run("  •  ".join(contact_parts))
        run_contact.font.name = 'Calibri'
        run_contact.font.size = Pt(9.5)
        run_contact.font.color.rgb = RGBColor(71, 85, 105)

    # 3. Section Rendering Function
    def add_section(heading_title: str, body_text: str):
        if not body_text.strip():
            return

        # Section Heading
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(10)
        p_head.paragraph_format.space_after = Pt(3)
        p_head.paragraph_format.keep_with_next = True
        run_head = p_head.add_run(heading_title.upper())
        run_head.font.name = 'Calibri'
        run_head.font.size = Pt(11)
        run_head.font.bold = True
        run_head.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy Blue

        # Content lines
        for raw_line in body_text.split("\n"):
            line = raw_line.strip()
            if not line:
                continue

            if line.startswith("•") or line.startswith("-") or line.startswith("*"):
                bullet_text = re.sub(r"^[-*•]\s*", "", line)
                p_bullet = doc.add_paragraph(style='List Bullet')
                p_bullet.paragraph_format.space_before = Pt(0)
                p_bullet.paragraph_format.space_after = Pt(2)
                p_bullet.paragraph_format.line_spacing = 1.15
                r = p_bullet.add_run(bullet_text)
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
            else:
                p_body = doc.add_paragraph()
                p_body.paragraph_format.space_before = Pt(2)
                p_body.paragraph_format.space_after = Pt(3)
                p_body.paragraph_format.line_spacing = 1.15
                r = p_body.add_run(line)
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                # Bold company/title lines if appropriate
                if any(kw in line.lower() for kw in ["inc.", "corp.", "technologies", "engineer", "lead", "architect", "scientist", "developer"]):
                    r.font.bold = True

    # Standard Section Flow
    section_map = {
        "summary": "Professional Summary",
        "skills": "Technical Proficiencies & Core Competencies",
        "experience": "Professional Experience",
        "projects": "Key Technical Projects",
        "education": "Education & Academic Credentials",
        "certifications": "Certifications & Accreditations",
        "awards": "Honors & Achievements"
    }

    seen = set()
    for sec_key, title in section_map.items():
        if sec_key in sections_dict and sections_dict[sec_key].strip():
            seen.add(sec_key)
            add_section(title, sections_dict[sec_key])

    # Custom sections
    for sec_key, content in sections_dict.items():
        if sec_key not in seen and sec_key != "header" and content.strip():
            add_section(sec_key.title(), content)

    # Save to byte stream
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io.getvalue()
