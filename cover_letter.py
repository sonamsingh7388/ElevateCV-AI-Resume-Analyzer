"""
Cover Letter & Recruiter Pitch Generator.
Creates customized executive cover letters and LinkedIn recruiter outreach pitches
tailored to the target job description and candidate background, with DOCX export.
"""

import io
from typing import Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from modules.gemini_client import generate_gemini_cover_letter


def generate_cover_letter_and_pitch(
    candidate_name: str,
    resume_text: str,
    jd_text: str,
    company: str = "Hiring Team",
    job_title: str = "Target Position",
    api_key: Optional[str] = None
) -> Dict[str, Any]:
    """
    Generates tailored cover letter and LinkedIn pitch using Gemini or offline heuristic engine.
    """
    return generate_gemini_cover_letter(
        candidate_name=candidate_name,
        resume_text=resume_text,
        jd_text=jd_text,
        company=company,
        job_title=job_title,
        api_key=api_key
    )


def export_cover_letter_to_docx(
    candidate_name: str,
    contact_info: Dict[str, Any],
    company: str,
    job_title: str,
    letter_text: str
) -> bytes:
    """
    Builds a beautifully styled cover letter as a Microsoft Word document (.docx).
    """
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 41, 59)

    # Candidate Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    r_name = p_name.add_run(candidate_name.upper())
    r_name.font.size = Pt(16)
    r_name.font.bold = True
    r_name.font.color.rgb = RGBColor(15, 23, 42)

    # Contact line
    contacts = []
    if contact_info.get("email"):
        contacts.append(contact_info["email"])
    if contact_info.get("phone"):
        contacts.append(contact_info["phone"])
    if contact_info.get("linkedin"):
        contacts.append(contact_info["linkedin"])

    if contacts:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_after = Pt(18)
        r_c = p_c.add_run("  |  ".join(contacts))
        r_c.font.size = Pt(9.5)
        r_c.font.color.rgb = RGBColor(100, 116, 139)

    # Letter Body
    # Filter out markdown headers if present
    clean_lines = []
    for line in letter_text.split("\n"):
        if line.startswith("#") or "Part 2" in line or "LinkedIn" in line:
            if "Part 2" in line or "LinkedIn" in line:
                break
            continue
        clean_lines.append(line)

    clean_body = "\n".join(clean_lines).strip()

    for para in clean_body.split("\n\n"):
        para_clean = para.strip()
        if not para_clean:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        
        # Check bullet inside cover letter
        if para_clean.startswith("-") or para_clean.startswith("•"):
            for bullet in para_clean.split("\n"):
                b_clean = bullet.strip().lstrip("-•*").strip()
                p_b = doc.add_paragraph(style='List Bullet')
                p_b.paragraph_format.space_after = Pt(3)
                p_b.add_run(b_clean)
        else:
            p.add_run(para_clean.replace("**", ""))

    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    return doc_io.getvalue()
